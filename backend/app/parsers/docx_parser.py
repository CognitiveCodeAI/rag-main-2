"""DOCX document parser using python-docx."""

import hashlib
import io
from typing import Any

from docx import Document as DocxDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.table import Table

from .base import (
    Block,
    DocumentIR,
    DocumentParser,
    Page,
    ParseQuality,
    Section,
)


class DocxParser(DocumentParser):
    """Parser for DOCX documents."""
    
    @property
    def supported_types(self) -> list[str]:
        return ["docx"]
    
    def parse(self, raw_bytes: bytes, doc_ref: dict[str, Any]) -> DocumentIR:
        """Parse DOCX into DocumentIR."""
        doc = DocxDocument(io.BytesIO(raw_bytes))
        
        blocks: list[Block] = []
        sections: list[Section] = []
        section_stack: list[Section] = []
        total_offset = 0
        block_count = 0
        title = None
        
        # Process document body elements
        for element in doc.element.body:
            tag = element.tag.split("}")[-1] if "}" in element.tag else element.tag
            
            if tag == "p":
                # Paragraph
                para = None
                for p in doc.paragraphs:
                    if p._element is element:
                        para = p
                        break
                
                if para is None or not para.text.strip():
                    continue
                
                text = para.text.strip()
                start_offset = total_offset
                end_offset = total_offset + len(text)
                total_offset = end_offset + 1
                
                # Check for heading style
                style_name = para.style.name if para.style else ""
                is_heading = style_name.startswith("Heading")
                heading_level = None
                
                if is_heading:
                    try:
                        heading_level = int(style_name.replace("Heading ", ""))
                    except ValueError:
                        heading_level = 1
                    
                    block = Block(
                        type="heading",
                        text=text,
                        level=heading_level,
                        start_offset=start_offset,
                        end_offset=end_offset,
                    )
                    
                    # Update section tracking
                    section_id = hashlib.md5(
                        f"{block_count}:{text[:50]}".encode()
                    ).hexdigest()[:8]
                    
                    while section_stack and section_stack[-1].level >= heading_level:
                        section_stack.pop()
                    
                    section_path = [s.heading for s in section_stack] + [text]
                    
                    section = Section(
                        section_id=section_id,
                        heading=text,
                        level=heading_level,
                        block_refs=[block_count],
                        section_path=section_path,
                    )
                    sections.append(section)
                    section_stack.append(section)
                    
                    if heading_level == 1 and title is None:
                        title = text
                else:
                    # Check for list
                    is_list = para._element.pPr is not None and para._element.pPr.numPr is not None
                    
                    if is_list:
                        block = Block(
                            type="list",
                            list_items=[text],
                            start_offset=start_offset,
                            end_offset=end_offset,
                        )
                    else:
                        block = Block(
                            type="paragraph",
                            text=text,
                            start_offset=start_offset,
                            end_offset=end_offset,
                        )
                    
                    if section_stack:
                        section_stack[-1].block_refs.append(block_count)
                
                blocks.append(block)
                block_count += 1
            
            elif tag == "tbl":
                # Table
                table = None
                for t in doc.tables:
                    if t._element is element:
                        table = t
                        break
                
                if table is None:
                    continue
                
                cells: list[list[str]] = []
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells]
                    cells.append(row_cells)
                
                table_text = "\n".join(["\t".join(row) for row in cells])
                start_offset = total_offset
                end_offset = total_offset + len(table_text)
                total_offset = end_offset + 1
                
                block = Block(
                    type="table",
                    cells=cells,
                    start_offset=start_offset,
                    end_offset=end_offset,
                )
                
                if section_stack:
                    section_stack[-1].block_refs.append(block_count)
                
                blocks.append(block)
                block_count += 1
        
        # DOCX is single "page" - we treat it as one logical page
        pages = [Page(page_num=1, blocks=blocks)]
        
        # Extract title from core properties if not found
        if title is None and doc.core_properties.title:
            title = doc.core_properties.title
        
        quality = ParseQuality(
            parse_confidence=0.95,  # DOCX is well-structured
            ocr_used=False,
        )
        
        return DocumentIR(
            doc_ref=doc_ref,
            title=title,
            language=doc.core_properties.language,
            pages=pages,
            sections=sections,
            quality=quality,
        )
